"""Translate CBO NIPA projection workbooks (.xlsx) into long, machine-readable CSV.

Usage: python convert_nipa_to_csv.py [source] [output.csv] (optional): output file path. The folder must already exist. If omitted, the CSV files are written to the repository folder.
`source` may be a URL or a local .xlsx path. With no source, the workbook is
downloaded from the CBO publication page (PUBLICATION_URL).
"""
import io
import re
import sys
from pathlib import Path
from urllib.parse import urljoin, urlparse, unquote
from urllib.request import Request, urlopen

import pandas as pd
from openpyxl import load_workbook

try:
    from openpyxl.cell.rich_text import CellRichText, TextBlock
except ImportError:  # openpyxl < 3.1
    CellRichText = TextBlock = None

REPORT_NUMBER = 62876
PUBLICATION_URL = f"https://www.cbo.gov/publication/{REPORT_NUMBER}"
OUT_DIR = Path(__file__).parent
HEADERS = {"User-Agent": "Mozilla/5.0 (NIPA_data converter)"}

YEAR_RE = re.compile(r"(?:(actual)\D*)?(\d{4})", re.I)
TOTAL_COL_RE = re.compile(r"total|\d{4}\s*[-–—]\s*\d{2,4}", re.I)
NOTE_RE = re.compile(r"^(?:([a-z])\.\s+|notes?\b|data sources?\b|sources?\b)", re.I)
LETTER_NOTE_RE = re.compile(r"^([a-z])\.\s+(.*)", re.I | re.S)
GENERIC_NOTE_RE = re.compile(r"^(notes?|data sources?|sources?)\s*:?\s*(.*)", re.I | re.S)
NAV_RE = re.compile(r"\s*back to (the )?table of contents\.?\s*", re.I)


def split_label(value):
    """Return (clean_label, footnote_markers) using superscript formatting.

    Curly apostrophes are normalized to straight ones for consistent matching.
    """
    if CellRichText is not None and isinstance(value, CellRichText):
        text, marks = "", ""
        for part in value:
            if isinstance(part, TextBlock):
                if part.font is not None and part.font.vertAlign == "superscript":
                    marks += part.text.strip()
                else:
                    text += part.text
            else:
                text += str(part)
        return " ".join(text.replace("’", "'").split()), marks
    return " ".join(str(value).replace("’", "'").split()), ""


def _get(url):
    with urlopen(Request(url, headers=HEADERS), timeout=60) as resp:
        return resp.read()


def find_xlsx_url(page_url=PUBLICATION_URL):
    """Return the absolute URL of the first .xlsx linked from the publication page."""
    html = _get(page_url).decode("utf-8", errors="replace")
    links = re.findall(r'href=["\']([^"\']+\.xlsx)(?:\?[^"\']*)?["\']', html, re.I)
    if not links:
        sys.exit(f"No .xlsx link found on {page_url}")
    return urljoin(page_url, links[0])


def load_source(src=None):
    """Return (workbook_bytes, file_name) from a URL, a local path, or the CBO page."""
    if src and not src.lower().startswith(("http://", "https://")):
        p = Path(src)
        return io.BytesIO(p.read_bytes()), p.name
    url = src if src and src.lower().endswith(".xlsx") else find_xlsx_url(src or PUBLICATION_URL)
    print(f"Downloading {url}")
    return io.BytesIO(_get(url)), unquote(Path(urlparse(url).path).name)


def parse_year(cell):
    """Return (year, status) for a single-year header, else None."""
    if cell is None:
        return None
    s = str(cell)
    if TOTAL_COL_RE.search(s):  # cumulative columns, e.g. "Total, 2026–2036"
        return None
    m = YEAR_RE.search(s)
    if not m or not 1900 <= int(m[2]) <= 2200:
        return None
    return int(m[2]), "actual" if m[1] else "projected"


def merged_value(ws, row, col):
    """Value of a cell, resolving merged ranges (1-based row/col)."""
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return ws.cell(rng.min_row, rng.min_col).value
    return ws.cell(row, col).value


def _norm(s):
    """Normalize a label for header matching (lowercase, letters only, no plural)."""
    return re.sub(r"[^a-z ]", "", s.lower()).strip().rstrip("s")


def is_number(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def is_blank(cell):
    return cell.value is None or not str(cell.value).strip()


def parse_notes(rows, start, source, sheet, table_no):
    """Parse the notes block into row-scoped (lettered) and table-scoped notes."""
    notes = []
    for r in rows[start:]:
        if is_blank(r[0]):
            continue
        text = NAV_RE.sub(" ", split_label(r[0].value)[0]).strip()
        if not text:  # navigation-only cell (e.g. hyperlink to Contents)
            continue
        if m := LETTER_NOTE_RE.match(text):
            notes.append({"scope": "row", "marker": m[1].lower(),
                          "note_type": "footnote", "text": m[2].strip()})
        elif m := GENERIC_NOTE_RE.match(text):
            notes.append({"scope": "table", "marker": "",
                          "note_type": m[1].lower().rstrip("s"),
                          "text": m[2].strip()})
        elif notes:  # wrapped continuation of previous note
            notes[-1]["text"] = f"{notes[-1]['text']} {text}".strip()
        else:
            notes.append({"scope": "table", "marker": "",
                          "note_type": "note", "text": text})
    for n in notes:
        n.update(source=source, sheet=sheet, table=table_no)
    return notes


def infer_marker(label, markers):
    """Fallback: strip trailing footnote letter(s) not formatted as superscript."""
    if (len(label) > 3 and label[-1] in markers
            and label[-2].isalpha() and label[-2].islower()):
        return label[:-1], label[-1]
    return label, ""


def assign_sections(rows, start, end):
    """Split the body into runs separated by blank rows and name each run.

    A run ending in "Total X" is named "X"; otherwise it takes its first label.
    """
    section_of, run = {}, []

    def close():
        if not run:
            return
        labels = [split_label(rows[i][0].value)[0] for i in run]
        last = labels[-1]
        name = (last[5:].strip(" ,") if last.lower().startswith("total")
                else labels[0])
        name = name[:1].upper() + name[1:]
        for i in run:
            section_of[i] = name

    for i in range(start, end):
        if is_blank(rows[i][0]):
            close()
            run = []
        else:
            run.append(i)
    close()
    return section_of


def flatten_sheet(ws, source):
    rows = list(ws.iter_rows())
    vals = [[c.value for c in r] for r in rows]

    # Title / table number / units
    title_idx = next((i for i, r in enumerate(vals)
                      if r and isinstance(r[0], str)
                      and r[0].strip().lower().startswith("table")), None)
    if title_idx is None:
        return pd.DataFrame(), pd.DataFrame()
    title = " ".join(split_label(rows[title_idx][0].value)[0].split())
    m = re.match(r"Table\s+(\w+)\.?\s*(.*)", title)
    table_no, table_title = (m[1], m[2]) if m else ("", title)
    units = next((str(r[0]).strip() for r in vals[title_idx + 1:title_idx + 4]
                  if r and r[0] and str(r[0]).strip()), "")

    # Year header row (single-year columns only; cumulative totals skipped)
    hdr_idx = next(i for i, r in enumerate(vals)
                   if sum(parse_year(c) is not None for c in r[1:]) >= 3)
    year_cols = {j: parse_year(c) for j, c in enumerate(vals[hdr_idx])
                 if j > 0 and parse_year(c)}

    # "Actual" may sit in (possibly merged) header rows above the year row
    for j, (yr, st) in list(year_cols.items()):
        above = " ".join(str(merged_value(ws, i + 1, j + 1) or "")
                         for i in range(title_idx + 1, hdr_idx))
        if "actual" in above.lower():
            year_cols[j] = (yr, "actual")

    # Locate notes block: first note-like label with no numeric data
    notes_start = next(
        (i for i in range(hdr_idx + 1, len(rows))
         if rows[i][0].value and NOTE_RE.match(split_label(rows[i][0].value)[0])
         and not any(is_number(v) for v in vals[i][1:])),
        len(rows))
    notes = parse_notes(rows, notes_start, source, ws.title, table_no)
    for k, n in enumerate(notes, 1):
        n["note_id"] = (f"T{table_no}-{n['marker']}" if n["scope"] == "row"
                        else f"T{table_no}-N{k}")
    note_text = {n["marker"]: n["text"] for n in notes if n["scope"] == "row"}
    markers = set(note_text)
    applies = {m: [] for m in markers}

    section_of = assign_sections(rows, hdr_idx + 1, notes_start)

    records = []
    stack = []  # list of (label, marks) for open headers

    for r_i in range(hdr_idx + 1, notes_start):
        cell = rows[r_i][0]
        if is_blank(cell):
            stack.clear()  # blank row ends the block
            continue

        label, marks = split_label(cell.value)
        inferred = False
        if not marks:
            label, marks = infer_marker(label, markers)
            inferred = bool(marks)
        marks = marks.replace(",", "").lower()

        data = {j: vals[r_i][j] for j in year_cols
                if j < len(vals[r_i]) and is_number(vals[r_i][j])}

        low = label.lower()
        is_subtotal = low.startswith("subtotal")
        is_total = low.startswith("total")

        if not data:  # header row: its markers pass down to children
            stack.append((label, marks))
            continue

        # Find the header this subtotal/total summarizes
        target = (_norm(low.split(",", 1)[-1]) if is_subtotal
                  else _norm(label[5:]) if is_total else None)
        match_idx = None
        if target:
            for k in range(len(stack) - 1, -1, -1):
                if _norm(stack[k][0]) == target:
                    match_idx = k
                    break

        if is_subtotal and match_idx is not None:
            del stack[match_idx + 1:]  # parent = header X
        elif is_total:
            if match_idx is not None:
                del stack[match_idx:]  # parent = header X's parent
            else:
                stack.clear()  # unmatched total is top-level

        parent_path = " > ".join(s[0] for s in stack)

        series_id = f"T{table_no}-R{r_i + 1:03d}"
        row_links = [mk for mk in marks if mk in note_text]
        row_links += [mk for s in stack for mk in s[1]
                      if mk in note_text and mk not in row_links]
        for mk in row_links:
            applies[mk].append(series_id)
        fn_ids = ";".join(f"T{table_no}-{mk}" for mk in row_links)
        fn_text = "; ".join(f"{mk}. {note_text[mk]}" for mk in row_links)

        for j, v in data.items():
            year, status = year_cols[j]
            records.append({
                "source": source, "sheet": ws.title, "table": table_no,
                "series_id": series_id,
                "table_title": table_title, "units": units,
                "row": r_i + 1, "section": section_of.get(r_i, ""),
                "parent_path": parent_path,
                "item": label, "footnote": marks, "footnote_ids": fn_ids,
                "footnote_text": fn_text, "footnote_inferred": inferred,
                "is_subtotal": is_subtotal or is_total,
                "year": year, "status": status, "value": round(float(v), 3),
            })

        if is_subtotal and match_idx is not None:
            stack.pop()  # close header X after its subtotal

    for n in notes:
        if n["scope"] == "row":
            n["applies_to"] = " | ".join(applies.get(n["marker"], [])) or "unmatched"
    return pd.DataFrame(records), pd.DataFrame(notes)


def workbook_notes(wb, source):
    """Generic notes from non-table sheets (e.g. Contents) apply to everything."""
    out = []
    for ws in wb.worksheets:
        if ws.title.lower().startswith("table"):
            continue
        for r in ws.iter_rows():
            text = split_label(r[0].value)[0] if r[0].value else ""
            if text.lower().startswith(("for more information", "note", "source")):
                out.append({"source": source, "sheet": ws.title, "table": "",
                            "note_id": f"WB-N{len(out) + 1}",
                            "scope": "workbook", "marker": "", "note_type": "note",
                            "text": text, "applies_to": "all series"})
    return pd.DataFrame(out)


def process(data_bytes, name: str, out: Path | None = None):
    wb = load_workbook(data_bytes, data_only=True, rich_text=CellRichText is not None)
    data, notes = [], [workbook_notes(wb, name)]
    for ws in wb.worksheets:
        if not ws.title.lower().startswith("table"):
            continue
        d, n = flatten_sheet(ws, name)
        data.append(d)
        notes.append(n)
    result = pd.concat(data, ignore_index=True)

    dup = result.duplicated(["series_id", "year"], keep=False)
    if dup.any():
        raise ValueError("Duplicate series_id+year:\n"
                         f"{result.loc[dup, ['series_id', 'year', 'value']]}")

    out = out or OUT_DIR / (Path(name).stem + "_flat.csv")
    result.to_csv(out, index=False)

    all_notes = pd.concat(notes, ignore_index=True)
    fn = all_notes[all_notes["scope"] == "row"].reindex(
        columns=["note_id", "source", "sheet", "table", "marker", "applies_to", "text"])
    fn_out = out.with_name(out.stem.replace("_flat", "") + "_footnotes.csv")
    fn.to_csv(fn_out, index=False)

    print(f"{name}: {len(result):,} rows -> {out.name}; "
          f"{len(fn)} footnotes -> {fn_out.name}")
    if result["footnote_inferred"].any():
        print("  Inferred (non-superscript) markers - please verify:")
        for item in result.loc[result["footnote_inferred"], "item"].unique():
            print(f"    {item}")
    if (fn["applies_to"] == "unmatched").any():
        print("  Footnotes not matched to any row:",
              ", ".join(fn.loc[fn["applies_to"] == "unmatched", "marker"]))
    for t, grp in result.groupby("table"):
        print(f"  Table {t} sections:")
        for s in grp["section"].unique():
            print(f"    - {s}")


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else None
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else None
    process(*load_source(src), out)


if __name__ == "__main__":
    main()